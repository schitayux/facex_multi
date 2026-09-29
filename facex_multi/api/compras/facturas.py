"""
facex_multi.api.compras.facturas
--------------------------------
Factura de Compra (Purchase Invoice) desde FacEx Compras.

Factura directa (sin Orden ni Entrada previa): update_stock=1 cuando trae
al menos un producto de inventario, así la misma factura ingresa el stock.
Se graba con las validaciones nativas de ERPNext (antes se grababa con
ignore_validate/ignore_mandatory y los errores recién aparecían al validar).
"""
from __future__ import annotations

import json

import frappe
from frappe.utils import flt, today

from facex_multi.api.compras.common import (
    check_doc_access,
    default_uom,
    get_buying_price_list,
    get_conversion_rate,
    get_naming_series,
    get_payable_account,
    get_tax_rows,
    get_tax_templates_with_rate,
    is_own_scope,
    require_purchase,
    resolve_item_warehouse,
    resolve_purchase_tax_template,
)


# ---------------------------------------------------------------------------
# Búsquedas
# ---------------------------------------------------------------------------

@frappe.whitelist()
def search_items(txt: str = "", company: str = None) -> list:
    """Ítems comprables por código o nombre, con la bodega de recepción sugerida."""
    company = require_purchase(company, "puede_compras")

    results = frappe.db.get_all(
        "Item",
        filters=[["disabled", "=", 0], ["is_purchase_item", "=", 1]],
        or_filters=[["name", "like", f"%{txt}%"], ["item_name", "like", f"%{txt}%"]],
        fields=["name as item_code", "item_name", "item_group",
                "has_serial_no", "has_batch_no", "is_stock_item", "stock_uom"],
        order_by="item_code asc",
        limit=25,
    )

    from facex_multi.api.permissions import get_facex_allowed_warehouses
    allowed_warehouses = get_facex_allowed_warehouses(company, "compra")

    return [
        {
            "item_code":     r.item_code,
            "item_name":     r.item_name,
            "item_group":    r.item_group,
            "has_serial_no": int(r.has_serial_no or 0),
            "has_batch_no":  int(r.has_batch_no or 0),
            "is_stock_item": int(r.is_stock_item or 0),
            "uom":           r.stock_uom or default_uom(),
            "warehouse":     resolve_item_warehouse(r.item_code, company, allowed_warehouses) if r.is_stock_item else "",
        }
        for r in results
    ]


@frappe.whitelist()
def search_suppliers(txt: str = "", company: str = None) -> list:
    """Proveedores activos de la compañía (o sin compañía asignada)."""
    company = require_purchase(company, "puede_compras")
    results = frappe.db.sql(
        """
        SELECT name, supplier_name, tax_id
        FROM `tabSupplier`
        WHERE disabled = 0
          AND (supplier_name LIKE %(q)s OR name LIKE %(q)s OR tax_id LIKE %(q)s)
          AND (bfel_company = %(company)s OR bfel_company IS NULL OR bfel_company = '')
        ORDER BY supplier_name ASC
        LIMIT 20
        """,
        {"q": f"%{txt}%", "company": company},
        as_dict=True,
    )
    return [{"value": r.name, "label": r.supplier_name, "tax_id": r.tax_id or ""} for r in results]


# ---------------------------------------------------------------------------
# Lista / lectura
# ---------------------------------------------------------------------------

@frappe.whitelist()
def get_purchase_list(company: str = None, start_date: str = None, end_date: str = None,
                      supplier: str = None, docstatus: str = None, limit: int = 100) -> list:
    company = require_purchase(company, "puede_compras",
                               msg="No tiene permiso para consultar compras en FacEx.")

    filters = [["company", "=", company], ["is_return", "=", 0]]
    if start_date and end_date:
        filters.append(["posting_date", "between", [start_date, end_date]])
    elif start_date:
        filters.append(["posting_date", ">=", start_date])
    elif end_date:
        filters.append(["posting_date", "<=", end_date])
    if supplier:
        filters.append(["supplier", "=", supplier])
    if docstatus is not None and docstatus != "":
        filters.append(["docstatus", "=", int(docstatus)])
    if is_own_scope(company):
        filters.append(["owner", "=", frappe.session.user])

    return frappe.get_all(
        "Purchase Invoice",
        filters=filters,
        fields=["name", "supplier", "supplier_name", "posting_date", "bill_no", "bill_date",
                "grand_total", "outstanding_amount", "docstatus", "status", "currency", "owner"],
        order_by="posting_date desc, creation desc",
        limit=min(int(limit or 100), 500),
    )


@frappe.whitelist()
def get_purchase_invoice(name: str, company: str = None) -> dict:
    company = require_purchase(company, "puede_compras",
                               msg="No tiene permiso para ver compras en FacEx.")
    doc = frappe.get_doc("Purchase Invoice", name.strip())
    check_doc_access(doc, company)

    d = doc.as_dict()
    for item in d.get("items", []):
        meta = frappe.db.get_value(
            "Item", item.get("item_code"), ["has_serial_no", "has_batch_no", "is_stock_item"], as_dict=True
        ) or {}
        item["has_serial_no"] = int(meta.get("has_serial_no") or 0)
        item["has_batch_no"] = int(meta.get("has_batch_no") or 0)
        item["is_stock_item"] = int(meta.get("is_stock_item") or 0)
    d["owner_fullname"] = frappe.utils.get_fullname(doc.owner)
    return d


# ---------------------------------------------------------------------------
# Grabar / validar / cancelar
# ---------------------------------------------------------------------------

@frappe.whitelist()
def save_purchase_invoice(data_json: str) -> dict:
    """
    Crea o actualiza una Factura de Compra en borrador.
    data_json: {name?, company, supplier, posting_date, bill_no, bill_date,
                currency, tax_type (nombre real de la Plantilla de Impuestos de Compra),
                bfel_multi_tipo (Tipo FEL de encabezado),
                items: [{item_code, qty, rate, warehouse, serial_no, batch_no, bfel_multi_tipo}]}
    """
    data = json.loads(data_json) if isinstance(data_json, str) else data_json
    company = require_purchase(data.get("company"), "puede_compras",
                               msg="No tiene permiso para registrar compras en esta compañía.")
    name = (data.get("name") or "").strip()
    supplier = (data.get("supplier") or "").strip()
    if not supplier:
        frappe.throw("Seleccione el proveedor.")
    currency = data.get("currency") or frappe.db.get_value("Company", company, "default_currency") or "GTQ"
    posting_date = data.get("posting_date") or today()

    if name and frappe.db.exists("Purchase Invoice", name):
        doc = frappe.get_doc("Purchase Invoice", name)
        check_doc_access(doc, company)
        if doc.docstatus != 0:
            frappe.throw("Solo se pueden editar facturas en borrador.")
        doc.items = []
        doc.taxes = []
    else:
        doc = frappe.new_doc("Purchase Invoice")
        series = get_naming_series("Purchase Invoice", company)
        if series:
            doc.naming_series = series[0]

    doc.company             = company
    doc.supplier            = supplier
    doc.posting_date        = posting_date
    doc.set_posting_time    = 1
    doc.bill_no             = data.get("bill_no") or ""
    doc.bill_date           = data.get("bill_date") or posting_date
    doc.currency            = currency
    doc.conversion_rate     = get_conversion_rate(currency, company, posting_date)
    doc.buying_price_list   = get_buying_price_list()
    doc.price_list_currency = currency
    doc.plc_conversion_rate = doc.conversion_rate
    doc.ignore_pricing_rule = 1
    doc.taxes_and_charges   = resolve_purchase_tax_template(company, data.get("tax_type") or "")
    doc.credit_to           = get_payable_account(supplier, company, currency)
    doc.set_warehouse       = None
    if doc.meta.has_field("bfel_multi_tipo"):
        doc.bfel_multi_tipo = data.get("bfel_multi_tipo") or ""

    from facex_multi.api.permissions import get_facex_allowed_warehouses
    from facex_multi.api.stock import _resolve_batch
    allowed_warehouses = get_facex_allowed_warehouses(company, "compra")

    has_stock_item = False
    for row in data.get("items", []):
        item_code = (row.get("item_code") or "").strip()
        if not item_code:
            continue
        info = frappe.db.get_value(
            "Item", item_code, ["has_serial_no", "has_batch_no", "is_stock_item", "stock_uom"], as_dict=True
        )
        if not info:
            frappe.throw(f"Producto '{item_code}' no encontrado.")
        is_stock = int(info.is_stock_item or 0)
        has_serial = int(info.has_serial_no or 0)
        has_batch = int(info.has_batch_no or 0)
        stock_uom = info.stock_uom or default_uom()

        qty = flt(row.get("qty") or 1)
        rate = flt(row.get("rate") or 0)
        serial_no = (row.get("serial_no") or "").strip()
        batch_no = (row.get("batch_no") or "").strip()
        if has_serial and serial_no:
            serials = [s.strip() for s in serial_no.replace(",", "\n").split("\n") if s.strip()]
            qty = len(serials)
            serial_no = "\n".join(serials)

        wh = ""
        if is_stock:
            has_stock_item = True
            wh = row.get("warehouse") or resolve_item_warehouse(item_code, company, allowed_warehouses)
            if wh and allowed_warehouses is not None and wh not in allowed_warehouses:
                if row.get("warehouse"):
                    frappe.throw(f"La bodega '{wh}' no está habilitada para compra en su configuración.")
                wh = allowed_warehouses[0] if allowed_warehouses else ""
            if not wh:
                frappe.throw(f"Indique la bodega de recepción de {item_code}.")

        item_row = {
            "item_code":         item_code,
            "qty":               qty,
            "rate":              rate,
            "uom":               stock_uom,
            "stock_uom":         stock_uom,
            "conversion_factor": 1.0,
            "warehouse":         wh,
        }
        if row.get("bfel_multi_tipo"):
            item_row["bfel_multi_tipo"] = row.get("bfel_multi_tipo")
        if is_stock and has_serial and serial_no:
            item_row["use_serial_batch_fields"] = 1
            item_row["serial_no"] = serial_no
        if is_stock and has_batch:
            # Compra = recepción de mercadería: se auto-crea el maestro Batch si
            # el lote del proveedor aún no existe (mismo criterio que Entradas).
            item_row["use_serial_batch_fields"] = 1
            item_row["batch_no"] = _resolve_batch(item_code, batch_no, "in")
        doc.append("items", item_row)

    if not doc.items:
        frappe.throw("Agregue al menos un producto antes de guardar.")

    # Solo servicios/gastos → sin movimiento de inventario. Con update_stock=1
    # y ningún ítem de inventario ERPNext arma partidas en cero y rechaza la
    # validación ("Either debit or credit amount is required").
    doc.update_stock = 1 if has_stock_item else 0

    for tax_row in get_tax_rows(doc.taxes_and_charges):
        doc.append("taxes", tax_row)

    # ignore_permissions: el control es el permiso FacEx (require_purchase),
    # igual que antes; las validaciones nativas de ERPNext sí corren.
    doc.save(ignore_permissions=True)
    frappe.db.commit()
    return {"success": True, "name": doc.name}


@frappe.whitelist()
def submit_purchase_invoice(name: str) -> dict:
    doc = frappe.get_doc("Purchase Invoice", name.strip())
    company = require_purchase(doc.company, "puede_validar_compras",
                               msg="No tiene permiso para validar compras.")
    check_doc_access(doc, company)
    if doc.docstatus != 0:
        frappe.throw("La factura ya fue validada o cancelada.")
    doc.submit()
    frappe.db.commit()
    return {"success": True, "name": doc.name}


@frappe.whitelist()
def cancel_purchase_invoice(name: str) -> dict:
    doc = frappe.get_doc("Purchase Invoice", name.strip())
    company = require_purchase(doc.company, "puede_cancelar_compras",
                               msg="No tiene permiso para cancelar compras.")
    check_doc_access(doc, company)
    if doc.docstatus != 1:
        frappe.throw("Solo se pueden cancelar facturas validadas.")
    doc.cancel()
    frappe.db.commit()
    return {"success": True, "name": doc.name}


@frappe.whitelist()
def delete_purchase_invoice(name: str) -> dict:
    """Elimina un borrador (mismo permiso que grabarlo)."""
    doc = frappe.get_doc("Purchase Invoice", name.strip())
    company = require_purchase(doc.company, "puede_compras",
                               msg="No tiene permiso para eliminar compras.")
    check_doc_access(doc, company)
    if doc.docstatus != 0:
        frappe.throw("Solo se pueden eliminar facturas en borrador.")
    frappe.delete_doc("Purchase Invoice", doc.name, ignore_permissions=True)
    frappe.db.commit()
    return {"success": True}


# ---------------------------------------------------------------------------
# Carga desde Excel
# ---------------------------------------------------------------------------

@frappe.whitelist()
def process_purchase_excel(file_url: str, company: str = None) -> dict:
    """
    Parsea Excel con 2 hojas:
      Hoja 1 – ENCABEZADO (fila 2): proveedor | fecha_registro | no_factura | fecha_factura | moneda
      Hoja 2 – DETALLE (desde fila 2): codigo_item | precio_unitario | serie | lote
    Una fila por unidad para ítems con serie; varias filas sin serie se acumulan
    (agrupadas por código + lote).
    """
    import openpyxl
    from frappe.utils.file_manager import get_file_path

    company = require_purchase(company, "puede_compras")
    wb = openpyxl.load_workbook(get_file_path(file_url), data_only=True)

    from facex_multi.api.permissions import get_facex_allowed_warehouses
    allowed_warehouses = get_facex_allowed_warehouses(company, "compra")

    ws_h = wb.worksheets[0]
    rows_h = list(ws_h.iter_rows(min_row=2, values_only=True))
    if not rows_h or not rows_h[0][0]:
        frappe.throw("La hoja ENCABEZADO está vacía o mal formada.")
    h = rows_h[0]

    def _val(v, default=""):
        return str(v).strip() if v is not None else default

    header = {
        "supplier":     _val(h[0] if len(h) > 0 else ""),
        "posting_date": _val(h[1], today()),
        "bill_no":      _val(h[2] if len(h) > 2 else ""),
        "bill_date":    _val(h[3] if len(h) > 3 else None, _val(h[1], today())),
        "currency":     _val(h[4] if len(h) > 4 else "GTQ", "GTQ").upper() or "GTQ",
    }
    if len(wb.worksheets) < 2:
        frappe.throw("Falta la hoja DETALLE (segunda hoja del archivo).")
    rows_d = list(wb.worksheets[1].iter_rows(min_row=2, values_only=True))

    order, groups = [], {}
    for row in rows_d:
        if not row or not row[0]:
            continue
        item_code = str(row[0]).strip()
        rate = flt(row[1] if len(row) > 1 else 0)
        serial = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        batch = str(row[3]).strip() if len(row) > 3 and row[3] else ""
        key = (item_code, batch)
        if key not in groups:
            order.append(key)
            groups[key] = {"rate": rate, "serials": [], "qty": 0, "batch": batch}
        if serial:
            groups[key]["serials"].append(serial)
        else:
            groups[key]["qty"] += 1
        if rate:
            groups[key]["rate"] = rate

    items, errors = [], []
    for item_code, _batch in order:
        g = groups[(item_code, _batch)]
        info = frappe.db.get_value(
            "Item", item_code,
            ["item_name", "item_group", "has_serial_no", "has_batch_no", "is_stock_item", "stock_uom"],
            as_dict=True,
        )
        if not info:
            errors.append(f"Producto '{item_code}' no encontrado en el sistema.")
            continue
        has_serial = int(info.has_serial_no or 0)
        has_batch = int(info.has_batch_no or 0)
        if has_serial and g["serials"]:
            qty, serial_no = len(g["serials"]), "\n".join(g["serials"])
        else:
            qty, serial_no = g["qty"] or 1, ""
        if has_serial and not serial_no:
            errors.append(f"'{item_code}' maneja series pero no se encontraron números de serie.")
        if has_batch and not g["batch"]:
            errors.append(f"'{item_code}' se gestiona por lote pero la columna 'lote' viene vacía.")
        items.append({
            "item_code":     item_code,
            "item_name":     info.item_name,
            "item_group":    info.item_group,
            "has_serial_no": has_serial,
            "has_batch_no":  has_batch,
            "is_stock_item": int(info.is_stock_item or 0),
            "uom":           info.stock_uom or default_uom(),
            "qty":           qty,
            "rate":          g["rate"],
            "warehouse":     resolve_item_warehouse(item_code, company, allowed_warehouses) if info.is_stock_item else "",
            "serial_no":     serial_no,
            "batch_no":      g["batch"],
        })

    return {"header": header, "items": items, "errors": errors}


@frappe.whitelist()
def get_purchase_tax_templates(company: str = None) -> list:
    company = require_purchase(company, "puede_compras")
    return get_tax_templates_with_rate(company)
